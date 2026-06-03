"""XLSX draft extraction (deterministic, no LLM).

OOXML spreadsheets are zip containers of XML parts. The drafter:

- Pulls workbook frontmatter via direct zip + ElementTree parsing (`xl/workbook.xml`,
  `docProps/core.xml`, `docProps/app.xml`) — dep-light. A second openpyxl pass
  (`data_only=False`) counts formula cells per sheet; the zip member list surfaces
  external links, pivot tables, macros, and defined names.
- Emits one Section per worksheet, addressed `sheet=<name>`, holding one rendered
  segment (also `sheet=<name>`; a normalizer carving a sub-region appends
  `&bbox=<A1-range>`). Static-data sheets get a markdown table fingerprinted via
  simhash. Sheets whose own cells carry formulas or external references — plus
  chart-only and empty sheets — get a body-empty segment and an address-scoped
  `format-loss` issue (spec §4.3.3.1 shape; reconciliation #2 vs the reference's
  CarbonAi `format_loss`). Per-sheet *classification* stays the normalizer's job.

`openpyxl` (the `[office]` extra) is imported lazily, so `import corpus.draft` works
without it; drafting an actual xlsx without the extra raises a clear ImportError.

Mode: `body-draft`.
"""

from __future__ import annotations

import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from corpus import recordbuild, touches
from corpus.draft import DrafterResult, register
from corpus.fingerprint import algos_for_atom, text_fingerprints
from corpus.functional_uri import quote_value
from corpus.segments import Section, Segment

_NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    "dc": "http://purl.org/dc/elements/1.1/",
    "dcterms": "http://purl.org/dc/terms/",
    "ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties",
}

_ROW_CAP = 1000
_XLSX_SCHEMA_ID = "application/application_vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _openpyxl():
    try:
        import openpyxl
    except ImportError as e:
        raise ImportError(
            "xlsx drafting requires the `[office]` extra. Install with: "
            "uv pip install 'ath-corpus[office]'"
        ) from e
    return openpyxl


def loss_issue(
    *, address: str, severity: str, reason: str, detector: str, **fields: Any
) -> dict[str, Any]:
    """A spec §4.3.3.1 `format-loss` issue, segment-scoped via `address`."""
    return {
        "id": "format-loss",
        "severity": severity,
        "resolution": "open",
        "detector": detector,
        "address": address,
        "fields": {"reason": reason, **fields},
    }


def draft(
    xlsx_path: Path,
    *,
    build: recordbuild.Build,
    corpus_root: Path | None = None,
    record_id: str | None = None,
    record_metadata: dict[str, Any] | None = None,
    canonical_algo: str | None = None,
    fingerprint: bool | str | list[str] = False,
) -> DrafterResult:
    detector = touches.script_identifier("draft.xlsx")
    fields: dict[str, Any] = {}

    with zipfile.ZipFile(xlsx_path) as zf:
        members = set(zf.namelist())
        if "xl/workbook.xml" in members:
            sheet_names = _read_sheet_names(zf)
            if sheet_names:
                fields["sheet_names"] = sheet_names
                fields["sheet_count"] = len(sheet_names)
        if "docProps/core.xml" in members:
            fields.update(_read_core_props(zf))
        if "docProps/app.xml" in members:
            fields.update(_read_app_props(zf))
        complexity = _zip_complexity(zf, members)

    sheet_meta = _sheet_formula_metadata(xlsx_path)
    complexity["sheet_meta"] = sheet_meta
    complexity["formula_count"] = sum(m["formula_count"] for m in sheet_meta.values())

    fields["formula_count"] = complexity["formula_count"]
    fields["external_links"] = complexity["external_links"]
    if complexity["named_ranges"]:
        fields["named_ranges"] = complexity["named_ranges"]
    fields["has_pivot_tables"] = complexity["has_pivot_tables"]
    fields["has_macros"] = complexity["has_macros"]

    sections, issues = _draft_sheet_sections(
        xlsx_path, complexity, detector, algos_for_atom("text", fingerprint)
    )

    recordbuild.add_blocks(build, sections)
    return {
        "fields": fields,
        "embeds": [],
        "title": fields.get("workbook_title"),
        "issues": issues,
    }


register(_XLSX_SCHEMA_ID)(draft)


# ---------- per-sheet sections ---------- #


def _draft_sheet_sections(
    xlsx_path: Path, complexity: dict[str, Any], detector: str, text_algos: list[str]
) -> tuple[list[Section], list[dict[str, Any]]]:
    """One Section per worksheet (+ chart-sheet sections). Returns (sections, issues)
    where issues are address-scoped format-loss entries."""
    sheet_meta: dict[str, dict[str, Any]] = complexity.get("sheet_meta", {})
    wb = _openpyxl().load_workbook(xlsx_path, read_only=True, data_only=True)
    sections: list[Section] = []
    issues: list[dict[str, Any]] = []
    try:
        for ws in wb.worksheets:
            meta = sheet_meta.get(ws.title, {})
            seg, seg_issues = _worksheet_segment(
                ws,
                formula_count=meta.get("formula_count", 0),
                uses_external_ref=meta.get("uses_external_ref", False),
                detector=detector,
                text_algos=text_algos,
            )
            sections.append(_wrap_sheet_section(ws.title, seg))
            issues.extend(seg_issues)
        for cs in wb.chartsheets:
            seg, seg_issues = _chartsheet_segment(cs, detector)
            sections.append(_wrap_sheet_section(cs.title, seg))
            issues.extend(seg_issues)
    finally:
        wb.close()
    return sections, issues


def _wrap_sheet_section(sheet_title: str, child: Segment) -> Section:
    return Section(address=_segment_address(sheet_title), entry=sheet_title, segments=[child])


def _worksheet_segment(
    ws, *, formula_count: int, uses_external_ref: bool, detector: str, text_algos: list[str]
) -> tuple[Segment, list[dict[str, Any]]]:
    """Render a worksheet to a Segment + any address-scoped issues."""
    address = _segment_address(ws.title)

    if formula_count > 0 or uses_external_ref:
        return _empty_segment(address), [
            _formula_issue(ws.title, address, formula_count, uses_external_ref, detector)
        ]

    body, truncated = _render_worksheet(ws, _ROW_CAP)
    if not body.strip():
        return _empty_segment(address), [_empty_sheet_issue(ws.title, address, detector)]

    issues: list[dict[str, Any]] = []
    if truncated:
        reason = (
            f"sheet {ws.title!r} truncated to the first {_ROW_CAP} rows; "
            f"further rows omitted from the body"
        )
        issues.append(
            loss_issue(
                address=address,
                severity="warning",
                reason=reason,
                detector=detector,
                row_cap=_ROW_CAP,
            )
        )
    seg = Segment(
        atom="text",
        address=address,
        perceptual=text_fingerprints(body, text_algos),
        body=body,
    )
    return seg, issues


def _chartsheet_segment(cs, detector: str) -> tuple[Segment, list[dict[str, Any]]]:
    address = _segment_address(cs.title)
    issue = loss_issue(
        address=address,
        severity="warning",
        reason=f"chart-only sheet {cs.title!r}; no cell content to render",
        detector=detector,
    )
    return _empty_segment(address), [issue]


def _empty_segment(address: str) -> Segment:
    """A body-empty text segment (no fingerprint) — the lossless markdown table could
    not be produced (formula / external-ref / chart-only / empty sheet)."""
    return Segment(atom="text", address=address, perceptual=None, body="")


def _formula_issue(
    title: str, address: str, formula_count: int, uses_external_ref: bool, detector: str
) -> dict[str, Any]:
    parts = []
    if formula_count > 0:
        parts.append(f"{formula_count} formula cell(s)")
    if uses_external_ref:
        parts.append("external workbook references")
    composition = " and ".join(parts) if parts else "computed cells"
    return loss_issue(
        address=address,
        severity="warning",
        reason=(
            f"worksheet {title!r} contains {composition}; cells are computed snapshots "
            f"and the formulas / references are not preserved in the draft body"
        ),
        detector=detector,
        formula_count=formula_count,
        external_references=uses_external_ref,
    )


def _empty_sheet_issue(title: str, address: str, detector: str) -> dict[str, Any]:
    return loss_issue(
        address=address,
        severity="info",
        reason=f"worksheet {title!r} is empty — no cells with values",
        detector=detector,
    )


def _segment_address(sheet_title: str) -> str:
    """`sheet=<name>` — the worksheet axis. A normalizer carving a sub-region appends
    `&bbox=<A1-range>`; query-reserved chars in the name are percent-encoded."""
    return f"sheet={quote_value(sheet_title)}"


# ---------- worksheet → markdown ---------- #


def _rows_to_markdown(rows: list[list[str]], truncated: bool) -> tuple[str, bool]:
    """Shared worksheet-rows → markdown-table tail for the xlsx + xls drafters. `rows` are
    already `_fmt_cell`'d strings, with leading all-blank rows skipped and the row cap
    applied by the caller (those steps are iteration-specific). Trims trailing blank rows
    AND trailing all-blank columns, synthesizes a header (column letters where the first
    row is blank), and emits the table. Returns `(body, truncated)`."""
    while rows and not any(cell for cell in rows[-1]):
        rows.pop()
    if not rows:
        return ("", truncated)
    width = max(len(r) for r in rows)
    for r in rows:
        if len(r) < width:
            r.extend("" for _ in range(width - len(r)))
    last_col = max((i for r in rows for i, cell in enumerate(r) if cell), default=-1)
    if last_col < 0:
        return ("", truncated)
    width = last_col + 1
    rows = [r[:width] for r in rows]
    header = rows[0]
    if all(not cell for cell in header):
        header = [_col_label(i) for i in range(width)]
    else:
        header = [cell or _col_label(i) for i, cell in enumerate(header)]
    return (_emit_markdown_table(header, rows[1:]), truncated)


def _render_worksheet(ws, row_cap: int) -> tuple[str, bool]:
    """Stream cells via `iter_rows(values_only=True)` → markdown table. Returns
    `(body, truncated)`."""
    rows: list[list[str]] = []
    truncated = False
    for row_values in ws.iter_rows(values_only=True):
        if len(rows) >= row_cap:
            truncated = True
            break
        formatted = [_fmt_cell(v) for v in row_values]
        if not rows and not any(cell for cell in formatted):
            continue
        rows.append(formatted)
    return _rows_to_markdown(rows, truncated)


def _fmt_cell(value: Any) -> str:
    """Format one cell value as a markdown-table-safe string."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return str(value)
    s = str(value).replace("\\", "\\\\").replace("|", "\\|")
    return s.replace("\r\n", "<br>").replace("\n", "<br>").replace("\r", "<br>")


def _col_label(index: int) -> str:
    """Excel-style column label for a 0-indexed position (0 → 'Col A')."""
    n = index
    letters = ""
    while True:
        letters = chr(ord("A") + (n % 26)) + letters
        n = n // 26 - 1
        if n < 0:
            break
    return f"Col {letters}"


def _emit_markdown_table(header: list[str], data_rows: list[list[str]]) -> str:
    """Emit a GFM table with padded column widths (renderers ignore the padding)."""
    cols = len(header)
    widths = [len(h) for h in header]
    for row in data_rows:
        for i, cell in enumerate(row[:cols]):
            widths[i] = max(widths[i], len(cell))

    def line(cells: list[str]) -> str:
        return "| " + " | ".join(cells[i].ljust(widths[i]) for i in range(cols)) + " |"

    sep = "| " + " | ".join("-" * max(3, widths[i]) for i in range(cols)) + " |"
    out = [line(header), sep]
    out.extend(line(row[:cols] + [""] * (cols - len(row))) for row in data_rows)
    return "\n".join(out)


# ---------- complexity detection ---------- #


def _zip_complexity(zf: zipfile.ZipFile, members: set[str]) -> dict[str, Any]:
    return {
        "external_links": any(m.startswith("xl/externalLinks/") for m in members),
        "has_pivot_tables": any(m.startswith("xl/pivotTables/") for m in members),
        "has_macros": "xl/vbaProject.bin" in members,
        "named_ranges": _read_defined_names(zf) if "xl/workbook.xml" in members else [],
    }


def _read_defined_names(zf: zipfile.ZipFile) -> list[str]:
    try:
        with zf.open("xl/workbook.xml") as fh:
            tree = ET.parse(fh)
    except (ET.ParseError, KeyError):
        return []
    names_el = tree.getroot().find("main:definedNames", _NS)
    if names_el is None:
        return []
    return [dn.get("name") for dn in names_el.findall("main:definedName", _NS) if dn.get("name")]


_EXTERNAL_REF_RE = re.compile(r"\[\d+\]")


def _sheet_formula_metadata(xlsx_path: Path) -> dict[str, dict[str, Any]]:
    """Per-sheet formula count + whether any formula references an external workbook
    (a second openpyxl pass with `data_only=False`)."""
    out: dict[str, dict[str, Any]] = {}
    wb = _openpyxl().load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        for ws in wb.worksheets:
            count = 0
            uses_external = False
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    count += 1
                    if cell.value and _EXTERNAL_REF_RE.search(str(cell.value)):
                        uses_external = True
            out[ws.title] = {"formula_count": count, "uses_external_ref": uses_external}
    finally:
        wb.close()
    return out


# ---------- frontmatter extraction ---------- #


def _read_sheet_names(zf: zipfile.ZipFile) -> list[str]:
    try:
        with zf.open("xl/workbook.xml") as fh:
            tree = ET.parse(fh)
    except (ET.ParseError, KeyError):
        return []
    sheets_el = tree.getroot().find("main:sheets", _NS)
    if sheets_el is None:
        return []
    return [sh.get("name") for sh in sheets_el.findall("main:sheet", _NS) if sh.get("name")]


def _read_core_props(zf: zipfile.ZipFile) -> dict[str, Any]:
    try:
        with zf.open("docProps/core.xml") as fh:
            root = ET.parse(fh).getroot()
    except (ET.ParseError, KeyError):
        return {}
    out: dict[str, Any] = {}
    for qname, key in (
        ("dc:creator", "workbook_creator"),
        ("cp:lastModifiedBy", "workbook_modified_by"),
        ("dcterms:created", "workbook_created_at"),
        ("dcterms:modified", "workbook_modified_at"),
        ("dc:title", "workbook_title"),
    ):
        if v := _text(root, qname):
            out[key] = v
    return out


def _read_app_props(zf: zipfile.ZipFile) -> dict[str, Any]:
    try:
        with zf.open("docProps/app.xml") as fh:
            root = ET.parse(fh).getroot()
    except (ET.ParseError, KeyError):
        return {}
    out: dict[str, Any] = {}
    for el in root.iter():
        local = re.sub(r"^\{[^}]+\}", "", el.tag)
        if local == "Application" and el.text:
            out["workbook_application"] = el.text.strip()
        elif local == "Company" and el.text:
            out["workbook_company"] = el.text.strip()
    return out


def _text(root: ET.Element, qname: str) -> str:
    el = root.find(qname, _NS)
    return el.text.strip() if el is not None and el.text is not None else ""
